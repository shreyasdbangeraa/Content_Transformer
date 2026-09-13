import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single", inside_v=False):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    insideV = OxmlElement('w:insideV')
    insideV.set(qn('w:val'), 'single' if inside_v else 'none')
    if inside_v:
        insideV.set(qn('w:sz'), sz)
        insideV.set(qn('w:space'), '0')
        insideV.set(qn('w:color'), color)
    tblBorders.append(insideV)
    tblPr.append(tblBorders)

def add_page_number_to_run(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def generate_architecture_doc():
    doc = Document()

    # Configure A4 Page Setup with precise engineering margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4 Width (210 mm)
    section.page_height = Inches(11.69) # A4 Height (297 mm)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    # Header: ConteX AI | SIH Architecture
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun1 = hp.add_run("ConteX AI  |  ")
    hrun1.font.name = 'Calibri'
    hrun1.font.size = Pt(8.5)
    hrun1.font.bold = True
    hrun1.font.color.rgb = RGBColor(30, 58, 138) # Deep Navy
    hrun2 = hp.add_run("SIH Architecture Specification")
    hrun2.font.name = 'Calibri'
    hrun2.font.size = Pt(8.5)
    hrun2.font.color.rgb = RGBColor(100, 116, 139) # Slate

    # Footer: Smart India Hackathon (SIH) | Page X
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    frun1 = fp.add_run("Smart India Hackathon (SIH)  |  Page ")
    frun1.font.name = 'Calibri'
    frun1.font.size = Pt(8.5)
    frun1.font.color.rgb = RGBColor(100, 116, 139)
    add_page_number_to_run(frun1)

    # Base typography
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(9)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    # =========================================================================
    # PAGE 1: TITLE, OVERVIEW & MAIN ARCHITECTURE DIAGRAM
    # =========================================================================

    # Title Block
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(1)
    p_title.paragraph_format.keep_with_next = True
    
    r_kicker = p_title.add_run("ConteX AI  —  ")
    r_kicker.font.size = Pt(14)
    r_kicker.font.bold = True
    r_kicker.font.color.rgb = RGBColor(30, 58, 138) # Deep Navy
    
    r_main = p_title.add_run("System Architecture")
    r_main.font.size = Pt(14)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    p_sub.paragraph_format.keep_with_next = True
    r_sub = p_sub.add_run("AI-Powered Information Verification & Transformation Platform • Technical Specification")
    r_sub.font.size = Pt(9)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Architecture Overview Paragraph (3-4 lines)
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(0)
    p_intro.paragraph_format.space_after = Pt(6)
    p_intro.paragraph_format.line_spacing = 1.12
    r_intro = p_intro.add_run(
        "ConteX AI follows a controlled information-processing pipeline that transforms unstructured information into researched, "
        "verified, traceable knowledge and reusable content. The architecture separates ingestion, AI analysis, research and verification, "
        "knowledge synthesis, integrity, and content transformation into distinct processing stages."
    )
    r_intro.font.size = Pt(8.5)
    r_intro.font.color.rgb = RGBColor(51, 65, 85)

    # -------------------------------------------------------------
    # MAIN ARCHITECTURE DIAGRAM (Visual Centered Flow Layout)
    # -------------------------------------------------------------
    p_diag_head = doc.add_paragraph()
    p_diag_head.paragraph_format.space_before = Pt(2)
    p_diag_head.paragraph_format.space_after = Pt(3)
    p_diag_head.paragraph_format.keep_with_next = True
    r_dh = p_diag_head.add_run("Main System Architecture Pipeline")
    r_dh.font.size = Pt(10)
    r_dh.font.bold = True
    r_dh.font.color.rgb = RGBColor(30, 58, 138)

    # Table representing the end-to-end layered architecture flow
    # Columns: [Side Component / Sources] | [Central Pipeline Flow] | [Side Storage / Anchors]
    diag_rows = [
        # Row 0: User & Web Layer
        ("INPUT CHANNELS\n• Raw Documents (PDF, DOCX, TXT)\n• Scraped Web URLs\n• Operator Text Prompts", 
         "1. PRESENTATION LAYER (Web Application UI)\nResponsive Workspace • 1-Click Evaluation Demo • Interactive Studio", 
         "CLIENT RUNTIME\nNext.js 15 (React 19)\nTailwind CSS • TypeScript"),
        
        # Row 1: Ingestion & Security
        ("SECURITY SHIELD\n• SSRF Protection (RFC-1918 Block)\n• XML Prompt Injection Delimiter", 
         "2. DOCUMENT & DATA PROCESSING LAYER\nSpatial Coordinate Parsing • Multi-Profile Web Crawler • Content Normalizer", 
         "LOCAL FILE STORE\nOriginal Source Archives\nTemporary Upload Buffer"),
        
        # Row 2: AI Analysis & Extraction
        ("AI ENGINE POOL\n• Google Gemini / OpenAI GPT-4o\n• Local Ollama / Offline Mock Provider", 
         "3. AI ANALYSIS & DECOMPOSITION LAYER\nFactual Claim Extraction • Entity/Timeline Detection • Deterministic PII Masking", 
         "SENSITIVITY ENGINE\nRegex Masking for Emails,\nIPs, Hosts, & Credentials"),
        
        # Row 3: Research & Verification
        ("EXTERNAL RESEARCH\n• Multi-Engine Search (Bing, DDG)\n• Direct Official Entity Probing", 
         "4. RESEARCH & VERIFICATION ENGINE\nTargeted Query Generation • 8-Tier Source Ranking • Discrepancy/Conflict Detection", 
         "SOURCE HIERARCHY\nTier 1: Gov/CERT/NIST Standards\nTier 2-8: Portals, Academic, News"),
        
        # Row 4: Canonical Knowledge Layer
        ("PROVENANCE TRACKING\n• Primary Document Fact\n• External Verified Fact • Inference", 
         "5. CANONICAL KNOWLEDGE LAYER (Single Source of Truth)\nPersistent Normalized Fact Base • Structured Timeline • Empirical Metrics • Risks", 
         "DATA STORES\n• In-DB 384-dim Vector Store\n• SQLite / PostgreSQL Metadata"),
        
        # Row 5: Integrity & Content Transformation (Dual column split in center)
        ("GOVERNANCE CONTROLS\n• Mandatory Human Approval Gate\n• Conversational AI Refinement Loop", 
         "6. TRANSFORMATION & QUALITY VERIFICATION\nSimultaneous 7-Format Generation • Claim-to-Page Grounding • 8-Metric Quality Radar", 
         "BLOCKCHAIN ANCHOR\n• SHA-256 Hash Chaining\n• Smart Contract (Sepolia/Mock)"),
        
        # Row 6: Final Output & Distribution
        ("OUTPUT FORMATS\nExec Summary • Slides (.pptx) • Docx\nThreat Advisory • Social Threads", 
         "7. FINAL OUTPUT & SYNDICATION LAYER\nVerified Communication Artefacts • Native File Downloader • n8n Webhook Dispatcher", 
         "AUTOMATION\nn8n Webhook Syndication\nLinkedIn, X & Webhook APIs")
    ]

    tbl_arch = doc.add_table(rows=len(diag_rows), cols=3)
    tbl_arch.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_arch, color="CBD5E1", sz="4", val="single", inside_v=True)

    col_w = [Inches(1.8), Inches(3.31), Inches(1.8)]

    for idx, (left_text, center_text, right_text) in enumerate(diag_rows):
        row = tbl_arch.rows[idx]
        c_left, c_center, c_right = row.cells[0], row.cells[1], row.cells[2]
        c_left.width, c_center.width, c_right.width = col_w[0], col_w[1], col_w[2]

        # Backgrounds: Left/Right side components = light slate; Center = clean white / tinted
        set_cell_background(c_left, "F8FAFC")
        set_cell_background(c_right, "F8FAFC")
        
        if idx in [0, 4, 6]:
            set_cell_background(c_center, "EFF6FF") # Light blue accent for key stages
        else:
            set_cell_background(c_center, "FFFFFF")

        set_cell_margins(c_left, top=40, bottom=40, left=60, right=60)
        set_cell_margins(c_center, top=40, bottom=40, left=70, right=70)
        set_cell_margins(c_right, top=40, bottom=40, left=60, right=60)

        # Populate left
        p_l = c_left.paragraphs[0]
        p_l.paragraph_format.space_before, p_l.paragraph_format.space_after = Pt(0), Pt(0)
        p_l.paragraph_format.line_spacing = 1.05
        lines_l = left_text.split("\n")
        r_l_title = p_l.add_run(lines_l[0] + "\n")
        r_l_title.font.bold = True
        r_l_title.font.size = Pt(7.5)
        r_l_title.font.color.rgb = RGBColor(30, 58, 138)
        r_l_body = p_l.add_run("\n".join(lines_l[1:]))
        r_l_body.font.size = Pt(7)
        r_l_body.font.color.rgb = RGBColor(71, 85, 105)

        # Populate center
        p_c = c_center.paragraphs[0]
        p_c.paragraph_format.space_before, p_c.paragraph_format.space_after = Pt(0), Pt(0)
        p_c.paragraph_format.line_spacing = 1.08
        lines_c = center_text.split("\n")
        r_c_title = p_c.add_run(lines_c[0] + "\n")
        r_c_title.font.bold = True
        r_c_title.font.size = Pt(8)
        r_c_title.font.color.rgb = RGBColor(15, 23, 42)
        r_c_body = p_c.add_run("\n".join(lines_c[1:]))
        r_c_body.font.size = Pt(7.5)
        r_c_body.font.color.rgb = RGBColor(51, 65, 85)

        # Populate right
        p_r = c_right.paragraphs[0]
        p_r.paragraph_format.space_before, p_r.paragraph_format.space_after = Pt(0), Pt(0)
        p_r.paragraph_format.line_spacing = 1.05
        lines_r = right_text.split("\n")
        r_r_title = p_r.add_run(lines_r[0] + "\n")
        r_r_title.font.bold = True
        r_r_title.font.size = Pt(7.5)
        r_r_title.font.color.rgb = RGBColor(30, 58, 138)
        r_r_body = p_r.add_run("\n".join(lines_r[1:]))
        r_r_body.font.size = Pt(7)
        r_r_body.font.color.rgb = RGBColor(71, 85, 105)

    # -------------------------------------------------------------
    # ARCHITECTURE LAYERS BREAKDOWN (Compact 2-Column Summary)
    # -------------------------------------------------------------
    p_layers_head = doc.add_paragraph()
    p_layers_head.paragraph_format.space_before = Pt(6)
    p_layers_head.paragraph_format.space_after = Pt(2)
    p_layers_head.paragraph_format.keep_with_next = True
    r_lh = p_layers_head.add_run("System Processing Layers")
    r_lh.font.size = Pt(10)
    r_lh.font.bold = True
    r_lh.font.color.rgb = RGBColor(30, 58, 138)

    layers_data = [
        ("1. Presentation Layer", "Next.js 15 web interface for file upload, live research monitoring, fact-check inspection, AI editing, and output publishing."),
        ("2. Processing & Security Layer", "Spatial PDF parsing, DOCX extraction, web crawling, prompt injection delimiters, and SSRF private IP blocking."),
        ("3. AI Analysis Layer", "Extracts claims, entities, chronological events, statistics, and risks while performing deterministic PII masking."),
        ("4. Research & Verification Layer", "Queries external sources via Bing/DuckDuckGo, ranks evidence by 8 authority tiers, and detects cross-source discrepancies."),
        ("5. Knowledge Core & Storage Layer", "Stores immutable Single Source of Truth (SSOT) records with claim provenance and internal RAG vector embeddings."),
        ("6. Integrity & Security Layer", "SHA-256 cryptographic hash chaining and EVM smart contract anchoring (ContentIntegrityRegistry.sol) for non-repudiation."),
        ("7. Content Transformation Layer", "Concurrently synthesizes 7 formats (briefings, PPTX slides, advisories, social threads) governed by an operator approval gate.")
    ]

    tbl_lay = doc.add_table(rows=4, cols=2)
    tbl_lay.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_lay, color="E2E8F0", sz="4", val="single", inside_v=True)

    lay_col_w = [Inches(3.45), Inches(3.46)]
    for idx, (l_title, l_desc) in enumerate(layers_data):
        r_idx = idx // 2
        c_idx = idx % 2
        cell = tbl_lay.rows[r_idx].cells[c_idx]
        cell.width = lay_col_w[c_idx]
        set_cell_background(cell, "FAFAFA")
        set_cell_margins(cell, top=35, bottom=35, left=60, right=60)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(0), Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r_t = p.add_run(l_title + ": ")
        r_t.font.bold = True
        r_t.font.size = Pt(8)
        r_t.font.color.rgb = RGBColor(15, 23, 42)
        r_d = p.add_run(l_desc)
        r_d.font.size = Pt(7.5)
        r_d.font.color.rgb = RGBColor(51, 65, 85)

    # Empty 8th slot in row 3 cell 1: make it a summary badge
    cell_empty = tbl_lay.rows[3].cells[1]
    cell_empty.width = lay_col_w[1]
    set_cell_background(cell_empty, "EFF6FF")
    set_cell_margins(cell_empty, top=35, bottom=35, left=60, right=60)
    p_e = cell_empty.paragraphs[0]
    p_e.paragraph_format.space_before, p_e.paragraph_format.space_after = Pt(0), Pt(0)
    r_et = p_e.add_run("Governance Model: ")
    r_et.font.bold = True
    r_et.font.size = Pt(8)
    r_et.font.color.rgb = RGBColor(30, 58, 138)
    r_ed = p_e.add_run("Mandatory Human Approval Gate prevents unverified distribution; conversational AI refinement updates versions.")
    r_ed.font.size = Pt(7.5)
    r_ed.font.color.rgb = RGBColor(30, 58, 138)

    # =========================================================================
    # PAGE BREAK TO PAGE 2
    # =========================================================================
    doc.add_page_break()

    # =========================================================================
    # PAGE 2: TECH STACK, DATA FLOW, PRINCIPLES & RUNTIME ARCHITECTURE
    # =========================================================================

    p_p2_head = doc.add_paragraph()
    p_p2_head.paragraph_format.space_before = Pt(0)
    p_p2_head.paragraph_format.space_after = Pt(3)
    p_p2_head.paragraph_format.keep_with_next = True
    r_p2 = p_p2_head.add_run("Technology Stack & Component Mapping")
    r_p2.font.size = Pt(10.5)
    r_p2.font.bold = True
    r_p2.font.color.rgb = RGBColor(30, 58, 138)

    # Technology Stack Table
    tech_stack = [
        ("Web Application", "Next.js 15 (App Router), React 19, TypeScript, Tailwind CSS", "User interface, interactive fact-check inspector & studio"),
        ("Backend / API", "FastAPI (Python 3.11), Uvicorn, Pydantic v2, SQLAlchemy 2.0", "Asynchronous REST application services & pipeline orchestration"),
        ("AI Processing", "Google Gemini API, OpenAI GPT-4o-mini, Local Ollama, Mock AI", "Canonical document decomposition & structured factual extraction"),
        ("Document Processing", "PyMuPDF (PDF spatial parsing), python-docx (Word extraction)", "Spatial text extraction preserving page coordinates & citations"),
        ("Web Scraping & Security", "BeautifulSoup4, HTTPX, URL SSRF Filter, XML Prompt Sanitizer", "External content ingestion with private IP and injection defense"),
        ("Research Engine", "Multi-Engine Search (Bing, DuckDuckGo via HTTP), Direct Domain Resolver", "External evidence retrieval anchored on document entities"),
        ("Verification & Scoring", "8-Tier Source Classifier, Claim Matcher, Flesch Readability Radar", "Evidence validation, page citation matching & quality scoring"),
        ("Knowledge Storage", "SQLite (Zero-config content_transformer.db) / PostgreSQL", "Single Source of Truth (SSOT), conflict records & project persistence"),
        ("Vector Retrieval (RAG)", "In-DB 384-dimensional vector arrays with Python Cosine Similarity", "Semantic search across internal reference knowledge documents"),
        ("Integrity & Governance", "SHA-256 Hash Chaining, Solidity Contract (Ethereum Sepolia/Mock)", "Cryptographic audit trail, version tracking & non-repudiation"),
        ("Content Generation", "python-pptx (Widescreen slides), python-docx, Structured JSON", "Simultaneous multi-format content synthesis & export"),
        ("Publishing Automation", "n8n Webhook Integration (Social Media AI Publisher)", "Automated post-approval distribution to LinkedIn, X, and webhooks")
    ]

    tbl_tech = doc.add_table(rows=len(tech_stack) + 1, cols=3)
    tbl_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_tech, color="CBD5E1", sz="4", val="single")

    t_col_w = [Inches(1.65), Inches(2.85), Inches(2.41)]

    # Header Row
    hdr = tbl_tech.rows[0]
    hdr.cells[0].width, hdr.cells[1].width, hdr.cells[2].width = t_col_w[0], t_col_w[1], t_col_w[2]
    for c in hdr.cells:
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=45, bottom=45, left=70, right=70)

    p_h0 = hdr.cells[0].paragraphs[0]
    p_h0.paragraph_format.space_before, p_h0.paragraph_format.space_after = Pt(0), Pt(0)
    r_h0 = p_h0.add_run("Architecture Component")
    r_h0.font.bold = True
    r_h0.font.size = Pt(8)
    r_h0.font.color.rgb = RGBColor(255, 255, 255)

    p_h1 = hdr.cells[1].paragraphs[0]
    p_h1.paragraph_format.space_before, p_h1.paragraph_format.space_after = Pt(0), Pt(0)
    r_h1 = p_h1.add_run("Technology / Implementation")
    r_h1.font.bold = True
    r_h1.font.size = Pt(8)
    r_h1.font.color.rgb = RGBColor(255, 255, 255)

    p_h2 = hdr.cells[2].paragraphs[0]
    p_h2.paragraph_format.space_before, p_h2.paragraph_format.space_after = Pt(0), Pt(0)
    r_h2 = p_h2.add_run("Role / Responsibility")
    r_h2.font.bold = True
    r_h2.font.size = Pt(8)
    r_h2.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (comp, tech, role) in enumerate(tech_stack):
        row = tbl_tech.rows[idx + 1]
        c0, c1, c2 = row.cells[0], row.cells[1], row.cells[2]
        c0.width, c1.width, c2.width = t_col_w[0], t_col_w[1], t_col_w[2]

        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        set_cell_background(c0, bg)
        set_cell_background(c1, bg)
        set_cell_background(c2, bg)
        set_cell_margins(c0, top=30, bottom=30, left=70, right=70)
        set_cell_margins(c1, top=30, bottom=30, left=70, right=70)
        set_cell_margins(c2, top=30, bottom=30, left=70, right=70)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before, p0.paragraph_format.space_after = Pt(0), Pt(0)
        r0 = p0.add_run(comp)
        r0.font.bold = True
        r0.font.size = Pt(7.5)
        r0.font.color.rgb = RGBColor(15, 23, 42)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before, p1.paragraph_format.space_after = Pt(0), Pt(0)
        r1 = p1.add_run(tech)
        r1.font.size = Pt(7.5)
        r1.font.color.rgb = RGBColor(51, 65, 85)

        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_before, p2.paragraph_format.space_after = Pt(0), Pt(0)
        r2 = p2.add_run(role)
        r2.font.size = Pt(7.5)
        r2.font.color.rgb = RGBColor(71, 85, 105)

    # -------------------------------------------------------------
    # DATA & CONTROL FLOW
    # -------------------------------------------------------------
    p_flow_head = doc.add_paragraph()
    p_flow_head.paragraph_format.space_before = Pt(6)
    p_flow_head.paragraph_format.space_after = Pt(2)
    p_flow_head.paragraph_format.keep_with_next = True
    r_fh = p_flow_head.add_run("Data & Control Flow")
    r_fh.font.size = Pt(10)
    r_fh.font.bold = True
    r_fh.font.color.rgb = RGBColor(30, 58, 138)

    # Visual Flow Badge Box
    tbl_fb = doc.add_table(rows=1, cols=1)
    tbl_fb.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_fb = tbl_fb.rows[0].cells[0]
    c_fb.width = Inches(6.91)
    set_cell_background(c_fb, "F1F5F9")
    set_cell_margins(c_fb, top=40, bottom=40, left=100, right=100)
    p_fb = c_fb.paragraphs[0]
    p_fb.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fb.paragraph_format.space_before, p_fb.paragraph_format.space_after = Pt(0), Pt(0)
    r_flow_text = p_fb.add_run(
        "INPUT  →  PROCESS  →  ANALYZE  →  RESEARCH  →  VERIFY  →  STORE / SYNTHESIZE  →  TRANSFORM  →  VALIDATE  →  OUTPUT"
    )
    r_flow_text.font.bold = True
    r_flow_text.font.size = Pt(8)
    r_flow_text.font.color.rgb = RGBColor(30, 58, 138)

    p_flow_desc = doc.add_paragraph()
    p_flow_desc.paragraph_format.space_before = Pt(3)
    p_flow_desc.paragraph_format.space_after = Pt(5)
    p_flow_desc.paragraph_format.line_spacing = 1.12
    r_fd = p_flow_desc.add_run(
        "Throughout this pipeline, extracted claims maintain explicit citation links to source page coordinates. "
        "External research findings are bound to specific empirical statements with provenance tags, ensuring that all "
        "synthesized deliverables remain factually grounded and traceable to primary evidence."
    )
    r_fd.font.size = Pt(8)
    r_fd.font.color.rgb = RGBColor(51, 65, 85)

    # -------------------------------------------------------------
    # KEY ARCHITECTURAL PRINCIPLES
    # -------------------------------------------------------------
    p_princ_head = doc.add_paragraph()
    p_princ_head.paragraph_format.space_before = Pt(4)
    p_princ_head.paragraph_format.space_after = Pt(2)
    p_princ_head.paragraph_format.keep_with_next = True
    r_ph = p_princ_head.add_run("Key Architectural Principles")
    r_ph.font.size = Pt(10)
    r_ph.font.bold = True
    r_ph.font.color.rgb = RGBColor(30, 58, 138)

    principles = [
        ("Evidence-Grounded Processing", "Eliminates hallucinations by establishing a verified factual knowledge core before invoking generation engines."),
        ("Separation of Analysis & Verification", "Decouples initial information extraction from external research, enabling cross-source discrepancy detection."),
        ("Traceable Knowledge & Provenance", "Retains source page numbers, bounding coordinates, and confidence weights on every individual factual claim."),
        ("Modular Asynchronous Pipeline", "Constructed with asynchronous microservices allowing parallel research execution, multi-format transformation, and pluggable AI providers."),
        ("Human-Controlled Final Output", "Enforces a strict Human Approval Gate that prevents autonomous unreviewed distribution while enabling conversational refinement.")
    ]

    for p_title, p_desc in principles:
        p_pr = doc.add_paragraph()
        p_pr.paragraph_format.space_before = Pt(1)
        p_pr.paragraph_format.space_after = Pt(1)
        p_pr.paragraph_format.line_spacing = 1.08
        r_pt = p_pr.add_run(f"• {p_title}: ")
        r_pt.font.bold = True
        r_pt.font.size = Pt(8)
        r_pt.font.color.rgb = RGBColor(15, 23, 42)
        r_pd = p_pr.add_run(p_desc)
        r_pd.font.size = Pt(8)
        r_pd.font.color.rgb = RGBColor(51, 65, 85)

    # -------------------------------------------------------------
    # DEPLOYMENT / RUNTIME ARCHITECTURE
    # -------------------------------------------------------------
    p_dep_head = doc.add_paragraph()
    p_dep_head.paragraph_format.space_before = Pt(5)
    p_dep_head.paragraph_format.space_after = Pt(2)
    p_dep_head.paragraph_format.keep_with_next = True
    r_dh = p_dep_head.add_run("Deployment & Runtime Architecture")
    r_dh.font.size = Pt(10)
    r_dh.font.bold = True
    r_dh.font.color.rgb = RGBColor(30, 58, 138)

    p_dep_body = doc.add_paragraph()
    p_dep_body.paragraph_format.space_before = Pt(0)
    p_dep_body.paragraph_format.space_after = Pt(0)
    p_dep_body.paragraph_format.line_spacing = 1.12
    r_db = p_dep_body.add_run(
        "The system operates as a decoupled microservice architecture: the Next.js 15 frontend communicates via asynchronous "
        "REST APIs with the FastAPI backend. For official Smart India Hackathon (SIH) evaluation, the platform is designed for "
        "complete local zero-configuration execution using an embedded SQLite database (content_transformer.db) and a built-in offline "
        "Mock AI provider, requiring no external database servers or paid API keys. For production scaling, the platform supports "
        "containerized deployment via Docker Compose with remote PostgreSQL persistence and cloud LLM endpoints (Gemini / OpenAI / Ollama)."
    )
    r_db.font.size = Pt(8)
    r_db.font.color.rgb = RGBColor(51, 65, 85)

    output_path = r"f:\content_transformation\ConteX_AI_System_Architecture_SIH.docx"
    doc.save(output_path)
    print(f"SUCCESS: Saved architecture document at {output_path}")

if __name__ == "__main__":
    generate_architecture_doc()
